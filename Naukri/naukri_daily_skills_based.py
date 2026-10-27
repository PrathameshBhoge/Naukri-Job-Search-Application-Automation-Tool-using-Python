#!/usr/bin/env python3
"""
Naukri daily workflow:  login + resume refresh  ->  search & shortlist  ->  YOU approve  ->  apply.

  python naukri_daily.py               # full flow
  python naukri_daily.py --skip-refresh   # skip the resume update step
  python naukri_daily.py --dry-run     # shortlist + show jobs, never apply

Rules built in:
  * Locations: Pune, Mumbai, Hyderabad, Bengaluru, Remote (Delhi NCR never)
  * Matching: prioritize skills in the job description over exact job-title wording
  * Experience: jobs asking up to 2 yrs minimum
  * Nothing is applied until you approve it at the prompt
  * Hard cap of 15 applications per calendar day (counted across runs)
  * External sites are skipped; unanswered questions are skipped where the site permits
Uses naukri_state.json created by naukri_refresh.py (log in once with HEADLESS=0).
"""
import re, sys, csv, json, random, time, subprocess, datetime, urllib.parse
from pathlib import Path

HERE = Path(__file__).resolve().parent
STATE, SEEN, LOG_JSON = HERE / "naukri_state.json", HERE / "seen_jobs.json", HERE / "applied_counts.json"
PROFILE_JSON = HERE / "naukri_profile.json"
TODAY = datetime.date.today().isoformat()
RESULT_CSV = HERE / f"tracker_rows_{TODAY}.csv"
ISSUES_XLSX = HERE / f"naukri_manual_review_{TODAY}.xlsx"

# ===================== YOUR SETTINGS =====================
DAILY_CAP = 15
MY_EXP_YEARS = 1.6
MAX_MIN_EXPERIENCE = 2                  # skip jobs whose minimum experience > this
CITIES = ["pune", "mumbai", "hyderabad", "bengaluru"]
REMOTE_WORDS = ["remote", "work from home", "wfh"]
NCR_WORDS = ["delhi", "gurgaon", "gurugram", "noida", "ghaziabad", "faridabad", "ncr"]
KEYWORDS = ["aws cloud support engineer", "cloud support engineer", "cloud operations engineer",
            "junior devops engineer", "technical support engineer", "application support engineer"]
MAX_PAGES_PER_SEARCH = 1

TITLE_REJECT = ["senior", "sr.", "sr ", "lead", "manager", "architect", "principal", "head", "sales",
                "customer service", "customer support", "voice", "bpo", "chat support", "intern", "trainee",
                "java developer", "python developer", ".net developer", "full stack", "data engineer",
                "data scientist", "servicenow"]
# Skill-first matching: do not require a particular title if the job description matches your skills.
MIN_SKILL_MATCHES = 2
CORE = ["aws", "linux", "unix", "cloud", "devops", "sql", "postgresql", "mysql", "monitoring", "jenkins",
        "nginx", "zabbix", "grafana", "bash", "shell", "ec2", "ci/cd"]
# Skills you have actually used (weights). Docker/Terraform/Kubernetes/Ansible/Azure are 0 until you confirm hands-on use.
MY_SKILLS = {"aws": 3, "ec2": 2, "s3": 1, "cloudwatch": 2, "route 53": 1, "alb": 1, "iam": 1,
             "linux": 3, "bash": 2, "shell": 2, "jenkins": 3, "ci/cd": 2, "cicd": 2, "git": 1,
             "nginx": 2, "apache": 1, "postgresql": 2, "postgres": 2, "mysql": 1,
             "zabbix": 2, "grafana": 2, "kibana": 1, "monitoring": 1, "troubleshooting": 1,
             "production support": 2, "l2": 1, "sla": 1, "incident": 1, "migration": 1,
             "docker": 0, "terraform": 0, "ansible": 0}
MISSING_PENALTY = {"gcp": 2, "azure": 1, "java": 2, "sap": 3, "vmware": 2, "windows server": 2,
                   "oracle": 2, "splunk": 1, "kubernetes": 1, "terraform": 0}
MIN_SCORE = 4
# =========================================================


# ---------- matching logic ----------
def parse_exp(text):
    m = re.search(r"(\d+(?:\.\d+)?)\s*[-–]\s*(\d+(?:\.\d+)?)", text or "")
    if m: return float(m.group(1)), float(m.group(2))
    m = re.search(r"(\d+(?:\.\d+)?)", text or "")
    return (float(m.group(1)),) * 2 if m else (None, None)


def _has(blob, term): return re.search(r"(?<![a-z0-9])" + re.escape(term) + r"(?![a-z0-9])", blob) is not None


def evaluate(job):
    """Skill-first match: -> (score, matched_skills) or (None, reason)."""
    title, loc = job["title"].lower(), job["location"].lower()

    # Keep seniority and clearly unrelated-role exclusions, but do not reject
    # a job merely because its title lacks "cloud", "AWS", or "DevOps".
    if any(w in title for w in TITLE_REJECT):
        return None, "unsuitable title"

    if loc and not (any(c in loc for c in CITIES) or any(r in loc for r in REMOTE_WORDS)):
        return None, "location not preferred"

    lo, hi = parse_exp(job["experience"])
    if lo is not None and lo > MAX_MIN_EXPERIENCE:
        return None, "needs more experience"
    if hi is not None and hi < 1:
        return None, "fresher-only"

    blob = " ".join([job["title"], job["snippet"], " ".join(job["tags"])]).lower()
    hits = {s: w for s, w in MY_SKILLS.items() if w > 0 and _has(blob, s)}
    if len(hits) < MIN_SKILL_MATCHES:
        return None, "not enough matching skills"

    # Require a relevant IT operations / cloud / DevOps role signal.
    # This prevents unrelated jobs (e.g. accounts roles) passing due to
    # incidental AWS/MySQL mentions in tags or descriptions.
    relevant_title_terms = (
        "cloud", "aws", "devops", "sre", "site reliability", "infrastructure",
        "system engineer", "systems engineer", "linux administrator",
        "cloud operations", "platform engineer", "observability",
        "production support", "application support", "technical support",
        "it support", "operations engineer", "noc", "middleware",
        "release engineer", "build engineer"
    )
    relevant_description_terms = (
        "cloud support", "cloud operations", "infrastructure support",
        "production support", "application support", "incident management",
        "monitoring", "deployment", "ci/cd", "server administration",
        "linux administration", "troubleshooting", "aws", "ec2",
        "cloudwatch", "jenkins", "nginx"
    )
    title_is_relevant = any(term in title for term in relevant_title_terms)
    description_is_relevant = any(term in blob for term in relevant_description_terms)
    if not title_is_relevant and not description_is_relevant:
        return None, "not a target IT/cloud/operations role"

    score = sum(hits.values()) - sum(
        p for s, p in MISSING_PENALTY.items() if _has(blob, s)
    )
    if lo is not None and lo <= MY_EXP_YEARS <= (hi if hi is not None else 99):
        score += 2
    if any(r in loc for r in REMOTE_WORDS) or any(c in loc for c in CITIES):
        score += 1

    if score < MIN_SCORE:
        return None, "skill match score below threshold"

    return score, sorted(hits, key=hits.get, reverse=True)


def parse_selection(s, n):
    s = (s or "").strip().lower()
    if s in ("", "none", "n", "no", "q"): return []
    if s in ("all", "a", "y", "yes"): return list(range(1, n + 1))
    out = set()
    for part in re.split(r"[,\s]+", s):
        m = re.fullmatch(r"(\d+)-(\d+)", part)
        if m: out.update(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit(): out.add(int(part))
    return sorted(i for i in out if 1 <= i <= n)


# ---------- daily cap ----------
def applied_today():
    try: return json.loads(LOG_JSON.read_text()).get(TODAY, 0)
    except Exception: return 0


def bump_applied(k=1):
    d = {}
    try: d = json.loads(LOG_JSON.read_text())
    except Exception: pass
    d[TODAY] = d.get(TODAY, 0) + k; LOG_JSON.write_text(json.dumps(d))


# ---------- browser bits ----------
def slug(s): return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def search_urls(keyword, page=1):
    q = {"k": keyword, "experience": 1, "jobAge": 3}
    for city in CITIES:
        base = f"https://www.naukri.com/{slug(keyword)}-jobs-in-{city}" + (f"-{page}" if page > 1 else "")
        yield city, base + "?" + urllib.parse.urlencode({**q, "l": city})
    q["wfhType"] = 2   # remote filter (verify on first run)
    yield "remote", f"https://www.naukri.com/{slug(keyword)}-jobs" + (f"-{page}" if page > 1 else "") + "?" + urllib.parse.urlencode(q)


def pause(a=2.0, b=4.5): time.sleep(random.uniform(a, b))


def any_visible(page, pattern):
    loc = page.get_by_text(re.compile(pattern, re.I))
    try:
        for i in range(min(loc.count(), 6)):
            if loc.nth(i).is_visible(): return True
    except Exception: pass
    return False


def scrape(page, url):
    page.goto(url, wait_until="domcontentloaded"); pause()
    try: page.wait_for_selector(".srp-jobtuple-wrapper, .jobTuple", timeout=15000)
    except Exception: return []
    out = []
    for card in page.locator(".srp-jobtuple-wrapper, .jobTuple").all():
        def txt(sel):
            try: return card.locator(sel).first.inner_text(timeout=1500).strip()
            except Exception: return ""
        try:
            href = card.locator("a.title").first.get_attribute("href") or ""
            tags = [t.strip() for t in card.locator(".tags-gt li, .tags li").all_inner_texts()]
        except Exception: continue
        out.append({"title": txt("a.title"), "company": txt("a.comp-name"), "experience": txt(".exp-wrap, .expwdth"),
                    "location": txt(".loc-wrap, .locWdth"), "snippet": txt(".job-desc, .job-description"),
                    "tags": tags, "url": href})
    return out


def load_profile():
    """Read user-supplied profile values. Blank/None values are never guessed."""
    if not PROFILE_JSON.exists():
        raise SystemExit(
            f"Missing {PROFILE_JSON.name}. Create it from the provided template and fill in your real details."
        )
    try:
        profile = json.loads(PROFILE_JSON.read_text(encoding="utf-8"))
    except Exception as exc:
        raise SystemExit(f"Could not read {PROFILE_JSON.name}: {exc}")
    return profile if isinstance(profile, dict) else {}


def is_naukri_url(url):
    try:
        host = urllib.parse.urlparse(url).hostname or ""
        return host == "naukri.com" or host.endswith(".naukri.com")
    except Exception:
        return False


def _clean_label(text):
    return re.sub(r"\s+", " ", (text or "").replace("*", " ")).strip().lower()


def _value_for_label(label, profile):
    """Map visible field/question labels to explicit profile values only."""
    label = _clean_label(label)
    if not label:
        return None
    aliases = {
        "full_name": ("full name", "name"),
        "email": ("email", "e-mail"),
        "phone": ("phone", "mobile", "contact number", "telephone"),
        "current_location": ("current location", "location", "city"),
        "experience_years": ("total experience", "years of experience", "experience in years", "work experience"),
        "notice_period": ("notice period", "when can you join", "availability to join", "joining time"),
        "current_ctc": ("current ctc", "current salary", "current compensation"),
        "expected_ctc": ("expected ctc", "expected salary", "salary expectation", "expected compensation"),
        "linkedin_url": ("linkedin",),
        "portfolio_url": ("portfolio", "personal website"),
        "current_company": ("current company", "employer"),
        "education": ("highest qualification", "education", "qualification", "degree"),
        "work_authorization_india": ("authorized to work", "work authorization", "right to work"),
        "willing_to_relocate": ("relocate", "willing to move", "open to relocation"),
    }
    for key, phrases in aliases.items():
        if any(phrase in label for phrase in phrases):
            value = profile.get(key)
            if value not in (None, "", []):
                return value
    # Explicit custom answers can cover employer-specific questions.
    # A key must be a distinctive substring of the question label.
    for phrase, value in (profile.get("question_answers") or {}).items():
        if phrase and phrase.lower() in label and value not in (None, "", []):
            return value
    return None


def _set_control(control, value):
    """Fill a text/select/radio/checkbox control. Returns True on success."""
    try:
        tag = control.evaluate("el => el.tagName.toLowerCase()")
        typ = (control.get_attribute("type") or "").lower()
        value_str = str(value)
        if tag == "select":
            options = control.locator("option").all()
            for option in options:
                label = (option.inner_text() or "").strip()
                val = option.get_attribute("value") or ""
                if label.lower() == value_str.lower() or val.lower() == value_str.lower():
                    control.select_option(value=val)
                    return True
            return False
        if typ in ("radio", "checkbox"):
            if isinstance(value, bool):
                # For yes/no questions, choose a matching visible option in the same group.
                desired = "yes" if value else "no"
                group = control.locator("xpath=..")
                candidates = group.locator("input[type=radio], input[type=checkbox]")
                for candidate in candidates.all():
                    lab = ""
                    try:
                        lab = candidate.evaluate("el => (el.labels && [...el.labels].map(x=>x.innerText).join(' ')) || el.getAttribute('aria-label') || el.value")
                    except Exception:
                        pass
                    if desired in (lab or "").lower():
                        candidate.check()
                        return True
                return False
            return False
        control.fill(value_str)
        return True
    except Exception:
        return False


def fill_application_form(page, profile):
    """Fill fields backed by profile.json; report unknown required fields but don't abort early."""
    unknown_required = []
    # Capture labels and nearby question text; avoid guessing for unrecognized fields.
    controls = page.locator("input:not([type=hidden]):not([type=submit]):not([type=button]), textarea, select")
    for i in range(min(controls.count(), 100)):
        control = controls.nth(i)
        try:
            if not control.is_visible() or control.is_disabled():
                continue
            typ = (control.get_attribute("type") or "").lower()
            if typ in ("file", "password", "submit", "button", "reset"):
                continue
            label = ""
            ident = control.get_attribute("id")
            if ident:
                try:
                    label = page.locator(f'label[for="{ident}"]').first.inner_text(timeout=500)
                except Exception:
                    pass
            if not label:
                label = control.get_attribute("aria-label") or control.get_attribute("placeholder") or control.get_attribute("name") or ""
            if not label:
                try:
                    label = control.locator("xpath=ancestor::*[self::label or self::div or self::section][1]").inner_text(timeout=500)
                except Exception:
                    label = ""
            value = _value_for_label(label, profile)
            if value is not None:
                if not _set_control(control, value):
                    # A matching value we cannot safely set is treated as unresolved.
                    if control.get_attribute("required") is not None:
                        unknown_required.append(label or "unidentified required field")
            elif control.get_attribute("required") is not None:
                current = ""
                try:
                    current = control.input_value()
                except Exception:
                    pass
                if not current:
                    unknown_required.append(label or control.get_attribute("name") or "unidentified required field")
        except Exception:
            continue
    # Deduplicate while preserving order. Unknown questions are left blank; the
    # caller will still try the site's normal submit path. The website itself
    # may block submission if a skipped question is mandatory.
    unknown_required = list(dict.fromkeys(x for x in unknown_required if x))
    if unknown_required:
        print("    Skipping unanswered required question(s): " + ", ".join(unknown_required[:8]))
    return unknown_required


def _click_save_if_answered(page):
    """Click a visible Save control only when a visible question input has a value."""
    try:
        save = page.get_by_role("button", name=re.compile(r"^\s*save\s*$", re.I))
        if not (save.count() and save.first.is_visible() and save.first.is_enabled()):
            return False
        fields = page.locator('input:not([type="hidden"]):not([type="button"]), textarea, [contenteditable="true"]')
        has_value = False
        for i in range(min(fields.count(), 40)):
            field = fields.nth(i)
            try:
                if not field.is_visible():
                    continue
                value = field.input_value() if field.get_attribute("contenteditable") is None else field.inner_text()
                if value and value.strip():
                    has_value = True
                    break
            except Exception:
                continue
        if has_value:
            save.first.click()
            pause(1, 2)
            return True
    except Exception:
        pass
    return False


def apply_to_job(page, url, profile):
    """Fill known fields, leave unknown questions blank, and try normal submission.

    External redirects are never followed. If the site requires an answer, it may
    reject the submission; this code does not fabricate personal information.
    """
    page.goto(url, wait_until="domcontentloaded"); pause(2, 4)
    if not is_naukri_url(page.url):
        return "external_skipped"
    if "nlogin" in page.url.lower(): return "login_needed"
    if any_visible(page, r"^\s*already applied\s*$"): return "already_applied"

    # Some listings are external applications from the start.
    if page.locator("#company-site-button").count() or any_visible(page, r"apply on company site|apply on employer website"):
        return "external_skipped"

    btn = page.locator("#apply-button")
    if not (btn.count() and btn.first.is_visible()):
        btn = page.get_by_role("button", name=re.compile(r"^\s*apply\s*$", re.I))
    if not (btn.count() and btn.first.is_visible()):
        return "no_apply_button"

    old_url = page.url
    btn.first.click(); pause(2, 4)
    # If clicking Apply sends us off Naukri, do not interact with the destination.
    if not is_naukri_url(page.url):
        return "external_skipped"
    if any_visible(page, r"successfully applied|applied successfully|application (has been )?sent"):
        return "applied"
    if "nlogin" in page.url.lower():
        return "login_needed"

    # Fill known fields. Unknown questions remain blank; do not stop the entire
    # application attempt merely because an unknown field exists.
    unknown = fill_application_form(page, profile)
    _click_save_if_answered(page)

    if any_visible(page, r"successfully applied|applied successfully|application (has been )?sent"):
        return "applied"

    # Submit only an explicit application-submit control, after the user approved this job batch.
    submit = page.get_by_role("button", name=re.compile(r"submit application|send application|submit your application|^submit$", re.I))
    if submit.count() and submit.first.is_visible() and submit.first.is_enabled():
        submit.first.click(); pause(2, 4)
        if not is_naukri_url(page.url):
            return "external_skipped"
        if any_visible(page, r"successfully applied|applied successfully|application (has been )?sent|application submitted"):
            return "applied"
        if "already applied" in page.locator("body").inner_text().lower():
            return "already_applied"
        # If the site kept us on the form, detect validation messages so the
        # tracker does not falsely claim success.
        try:
            body = page.locator("body").inner_text().lower()
            if any(term in body for term in ("required", "please answer", "mandatory", "please select", "invalid")):
                print("    Site blocked submission (likely a required unanswered question).")
                return "blocked_required_question"
        except Exception:
            pass
        return "unconfirmed"

    # No standard submit button found. Keep the result honest rather than
    # claiming success; a custom Save-only question panel may still need review.
    return "needs_manual"


# ---------- main flow ----------
def run_refresh():
    print("Step 1/4: login + resume refresh...")
    r = subprocess.run([sys.executable, str(HERE / "naukri_refresh.py")])
    if r.returncode != 0:
        ans = input("Refresh FAILED. Continue with shortlist anyway? [y/N]: ").strip().lower()
        if ans != "y": sys.exit("Stopped. Fix naukri_refresh.py first (see naukri_refresh.log).")


def shortlist(ctx, limit):
    seen = set(json.loads(SEEN.read_text())) if SEEN.exists() else set()
    page = ctx.new_page(); found, rejected = {}, {}
    print("Step 2/4: searching Naukri...")
    for kw in KEYWORDS:
        for pg in range(1, MAX_PAGES_PER_SEARCH + 1):
            for place, url in search_urls(kw, pg):
                print(f"  {kw} / {place}")
                for job in scrape(page, url):
                    key = job["url"].split("?")[0]
                    if not key or key in seen or key in found: continue
                    score, info = evaluate(job)
                    if score is None: rejected[info] = rejected.get(info, 0) + 1; continue
                    if score >= MIN_SCORE: job.update(score=score, matched=info, key=key); found[key] = job
                pause(3, 6)
    page.close()
    rows = sorted(found.values(), key=lambda j: -j["score"])[:limit]
    print(f"  Rejected: {rejected}")
    return rows, seen


def main():
    dry = "--dry-run" in sys.argv
    remaining = DAILY_CAP - applied_today()
    if remaining <= 0 and not dry:
        sys.exit(f"Daily cap of {DAILY_CAP} already reached today.")
    if "--skip-refresh" not in sys.argv and not dry: run_refresh()
    if not STATE.exists(): sys.exit("No saved session. Run naukri_refresh.py once with HEADLESS=0 first.")

    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(headless=False)
        ctx = b.new_context(storage_state=str(STATE), viewport={"width": 1366, "height": 850},
                            locale="en-IN", timezone_id="Asia/Kolkata")
        rows, seen = shortlist(ctx, min(DAILY_CAP, max(remaining, DAILY_CAP if dry else 0)))
        if not rows:
            print("No new matching jobs today."); b.close(); return

        print(f"\nStep 3/4: {len(rows)} matching jobs\n" + "-" * 78)
        for i, j in enumerate(rows, 1):
            print(f"{i:>2}. [{j['score']:>2}] {j['title']} | {j['company']} | {j['location']} | {j['experience']}")
            print(f"     skills: {', '.join(j['matched']) or '-'}\n     {j['url']}")
        if dry: print("\nDry run: nothing applied and jobs were not marked as seen."); b.close(); return

        picks = parse_selection(input(f"\nApprove which jobs? (e.g. 1,3,5-8 / all / none) [max {remaining} today]: "), len(rows))
        picks = picks[:remaining]
        if not picks: print("Nothing approved. Done."); b.close(); return
        # The job selection above is the approval. Do not ask for a second confirmation.
        print(f"Step 4/4: applying to {len(picks)} approved job(s)...")
        profile = load_profile()
        page = ctx.new_page(); out = []
        for n, idx in enumerate(picks, 1):
            j = rows[idx - 1]
            try: status = apply_to_job(page, j["url"], profile)
            except Exception as e: status = f"error: {str(e)[:60]}"
            print(f"  ({n}/{len(picks)}) {status:<16} {j['title']} @ {j['company']}")
            if status in ("applied", "unconfirmed"): bump_applied()
            if status not in ("external_skipped", "error"):
                seen.add(j["key"])
                SEEN.write_text(json.dumps(sorted(seen)))
            if status == "login_needed": print("  Session expired. Re-run naukri_refresh.py."); break
            out.append((j, status)); pause(8, 15)
        ctx.storage_state(path=str(STATE)); b.close()

    new = not RESULT_CSV.exists()
    with open(RESULT_CSV, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if new: w.writerow(["Company", "Role", "Location", "Job Link", "Source", "Resume Version",
                            "Date Applied", "Referral Contact", "Recruiter Contact", "Status"])
        for j, st in out:
            rv = "AWS/DevOps" if re.search(r"devops|infrastructure", j["title"].lower()) else "Cloud Support"
            done = st == "applied" or st == "already_applied"
            w.writerow([j["company"], j["title"], j["location"], j["url"], "Naukri", rv,
                        TODAY if done else "", "", "", st])

    # Keep external applications and cases needing manual attention in an Excel workbook.
    manual_statuses = {"external", "external_skipped", "needs_manual", "no_apply_button",
                       "blocked_required_question", "login_needed", "unconfirmed"}
    manual = [(j, st) for j, st in out if st in manual_statuses or st.startswith("error")]
    try:
        from openpyxl import Workbook, load_workbook
        if ISSUES_XLSX.exists():
            wb = load_workbook(ISSUES_XLSX)
            ws = wb["Needs Review"] if "Needs Review" in wb.sheetnames else wb.create_sheet("Needs Review")
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Needs Review"
            ws.append(["Date", "Status", "Company", "Role", "Location", "Job Link", "Notes"])
        existing = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) >= 6:
                existing.add((row[0], row[1], row[5]))
        for j, st in manual:
            key = (TODAY, st, j.get("url", ""))
            if key not in existing:
                note = "External application site skipped" if st == "external_skipped" else "Review application manually; submission not confirmed"
                ws.append([TODAY, st, j.get("company", ""), j.get("title", ""), j.get("location", ""), j.get("url", ""), note])
        wb.save(ISSUES_XLSX)
        print(f"Excel review sheet saved: {ISSUES_XLSX.name}")
    except ImportError:
        # Keep a CSV fallback if openpyxl is not installed.
        fallback = HERE / f"naukri_manual_review_{TODAY}.csv"
        is_new = not fallback.exists()
        with open(fallback, "a", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            if is_new: w.writerow(["Date", "Status", "Company", "Role", "Location", "Job Link", "Notes"])
            for j, st in manual:
                note = "External application site skipped" if st == "external_skipped" else "Review application manually; submission not confirmed"
                w.writerow([TODAY, st, j.get("company", ""), j.get("title", ""), j.get("location", ""), j.get("url", ""), note])
        print(f"openpyxl is not installed; CSV review sheet saved: {fallback.name}")

    print(f"\nApplied today: {applied_today()}/{DAILY_CAP}. Tracker saved to {RESULT_CSV.name}.")
    if manual:
        print("\nSkipped or requiring manual review (also saved to Excel):")
        for j, st in manual: print(f"  [{st}] {j['title']} @ {j['company']}\n     {j['url']}")
    unc = [j for j, st in out if st == "unconfirmed"]
    if unc: print(f"\n{len(unc)} application(s) could not be confirmed. Check 'Applied jobs' on Naukri.")


if __name__ == "__main__":
    main()
